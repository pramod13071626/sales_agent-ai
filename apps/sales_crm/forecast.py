"""Forecasting (apps/sales_crm/README.md §3, M2). Pure SQL + arithmetic, no AI requests.

Periods are fiscal quarters "FY2026-Q3" (fiscal year start month from crm_settings, January
by default; a FY is named after the calendar year it ends in). A deal belongs to a period by
its close date: `closed_at` for won deals, `expected_close` for open ones.

Categories (set per deal by the rep, forced by the DB trigger for won/lost):
    closed (won) · commit · best_case · pipeline · omitted (excluded)
Roll-ups follow the usual cumulative convention:
    Commit forecast    = closed + commit
    Best-case forecast = closed + commit + best_case
    Open pipeline      = commit + best_case + pipeline
    Weighted           = Σ open amount × probability (deal override, else the stage default)
All money is USD via fx_rates (deals.amount_usd, maintained by a trigger).

API /api/forecast: meta · (root) roll-up · changes · targets · fx · snapshot · export
"""

import threading
import time
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel, Field
from sqlalchemy import text

import auth
from apps.sales_crm import permissions
from db.connection import engine, get_session
from db.models.user import User

router = APIRouter(prefix="/api/forecast", tags=["CRM forecast"])
IST = ZoneInfo("Asia/Kolkata")
CATEGORIES = ["commit", "best_case", "pipeline", "omitted"]          # settable on open deals
CATEGORY_LABEL = {"closed": "Closed won", "commit": "Commit", "best_case": "Best case", "pipeline": "Pipeline", "omitted": "Omitted"}
OPEN_STAGES = ["intro", "discovery", "proposal", "pilot", "contract"]
STAGE_LABEL = {"intro": "Intro", "discovery": "Discovery", "proposal": "Proposal", "pilot": "Pilot", "contract": "Contract",
               "won": "Won", "lost": "Lost"}
READERS = ("sales_manager", "user", "viewer")


def _session():
    s = get_session()
    try:
        yield s
    finally:
        s.close()


def _q(s, sql: str, **kw) -> List[Dict[str, Any]]:
    return [dict(r) for r in s.execute(text(sql), kw).mappings()]


def _f(v) -> float:
    return float(v) if v is not None else 0.0


def _setting(s, key: str, default):
    v = s.execute(text("SELECT value FROM crm_settings WHERE key = :k"), {"k": key}).scalar()
    return default if v is None else v


_prob_cache: Tuple[float, Dict[str, int]] = (0.0, {})


def stage_probability(s) -> Dict[str, int]:
    """Stage default probabilities from crm_settings (cached 60 s; used per deal by the deals API)."""
    global _prob_cache
    if time.monotonic() - _prob_cache[0] < 60:
        return dict(_prob_cache[1])
    base = {"intro": 10, "discovery": 20, "proposal": 40, "pilot": 60, "contract": 80, "won": 100, "lost": 0}
    base.update(_setting(s, "stage_probability", {}) or {})
    _prob_cache = (time.monotonic(), base)
    return dict(base)


# ── Fiscal periods ────────────────────────────────────────────────────────────


def _add_months(d: date, n: int) -> date:
    m = d.month - 1 + n
    return date(d.year + m // 12, m % 12 + 1, 1)


def period_of(d: date, start_month: int) -> str:
    months_in = (d.month - start_month) % 12
    fy_start_year = d.year if d.month >= start_month else d.year - 1
    fy = fy_start_year if start_month == 1 else fy_start_year + 1
    return f"FY{fy}-Q{months_in // 3 + 1}"


def period_bounds(period: str, start_month: int) -> Tuple[date, date]:
    try:
        fy, q = int(period[2:6]), int(period[-1])
        assert period[6:8] == "-Q" and 1 <= q <= 4
    except (ValueError, AssertionError, IndexError):
        raise HTTPException(400, "period must look like FY2026-Q3")
    fy_start_year = fy if start_month == 1 else fy - 1
    start = _add_months(date(fy_start_year, start_month, 1), 3 * (q - 1))
    return start, _add_months(start, 3) - timedelta(days=1)


def period_list(start_month: int, today: Optional[date] = None, back: int = 2, ahead: int = 4) -> List[Dict[str, Any]]:
    today = today or datetime.now(IST).date()
    cur_start, _ = period_bounds(period_of(today, start_month), start_month)
    out = []
    for n in range(-back, ahead + 1):
        st = _add_months(cur_start, 3 * n)
        p = period_of(st, start_month)
        a, b = period_bounds(p, start_month)
        out.append({"key": p, "start": a, "end": b, "current": n == 0,
                    "label": f"{p.replace('-', ' ')} ({a.strftime('%b')}–{b.strftime('%b %Y')})"})
    return out


# ── Scope ─────────────────────────────────────────────────────────────────────


def visible_owners(s, user) -> List[Dict[str, Any]]:
    """Owners this user can pick in the owner filter / set targets for."""
    if permissions.is_admin(user):
        return _q(s, """SELECT id, coalesce(full_name, email) AS name, role FROM users
                        WHERE role NOT IN ('partner','viewer') AND is_active ORDER BY 2""")
    ids = [user.id] + (auth.team_user_ids(s, user.id) if user.role == "sales_manager" else [])
    return _q(s, "SELECT id, coalesce(full_name, email) AS name, role FROM users WHERE id = ANY(:ids) ORDER BY 2", ids=ids)


def can_edit_target(s, user, user_id: Optional[int], business_line_id: Optional[int]) -> bool:
    if permissions.is_admin(user):
        return True
    if user.role != "sales_manager" or business_line_id is not None:
        return False
    return user_id in auth.team_user_ids(s, user.id)


# ── Roll-up ───────────────────────────────────────────────────────────────────

DEAL_SQL = """
    SELECT d.id, d.name, d.account_id, a.display_name AS account_name, d.owner_user_id,
           coalesce(u.full_name, u.email, 'Unassigned') AS owner_name, d.business_line_id,
           coalesce(bl.name, 'No business line') AS business_line_name, d.stage, d.forecast_category,
           d.amount_usd, d.value_amount, d.currency, d.probability, d.expected_close,
           (d.closed_at AT TIME ZONE 'Asia/Kolkata')::date AS closed_on, d.next_step
    FROM deals d JOIN accounts a ON a.id = d.account_id
    LEFT JOIN users u ON u.id = d.owner_user_id LEFT JOIN business_lines bl ON bl.id = d.business_line_id
    WHERE d.account_id = ANY(:acl)
      AND (CAST(:bl AS int) IS NULL OR d.business_line_id = :bl)
      AND (CAST(:owner AS int) IS NULL OR d.owner_user_id = :owner)"""


def _bucket(d: Dict[str, Any]) -> str:
    if d["stage"] == "won":
        return "closed"
    if d["stage"] == "lost":
        return "lost"
    return d["forecast_category"] if d["forecast_category"] in CATEGORIES else "pipeline"


def _empty_totals() -> Dict[str, float]:
    return {k: 0.0 for k in ("closed", "commit", "best_case", "pipeline", "omitted", "weighted", "deals")}


def _finish(t: Dict[str, float], target: Optional[float]) -> Dict[str, Any]:
    t = {k: round(v, 2) for k, v in t.items()}
    t["deals"] = int(t["deals"])
    t["commit_forecast"] = round(t["closed"] + t["commit"], 2)
    t["best_case_forecast"] = round(t["closed"] + t["commit"] + t["best_case"], 2)
    t["open_pipeline"] = round(t["commit"] + t["best_case"] + t["pipeline"], 2)
    t["target"] = round(target, 2) if target is not None else None
    if target:
        gap = max(0.0, target - t["closed"])
        t["gap"] = round(gap, 2)
        t["attainment_pct"] = round(t["closed"] / target * 100, 1)
        t["coverage"] = round(t["open_pipeline"] / gap, 2) if gap else None
    else:
        t["gap"] = t["attainment_pct"] = t["coverage"] = None
    return t


def rollup(s, user, period: str, business_line_id: Optional[int] = None, owner_id: Optional[int] = None,
           group_by: str = "owner") -> Dict[str, Any]:
    if group_by not in ("owner", "business_line", "stage", "category"):
        raise HTTPException(400, "group_by must be owner, business_line, stage or category")
    sm = int(_setting(s, "fiscal_year_start_month", 1))
    start, end = period_bounds(period, sm)
    probs = stage_probability(s)
    acl = auth.account_scope(s, user)
    rows = _q(s, DEAL_SQL + """
          AND ((d.stage = 'won' AND (d.closed_at AT TIME ZONE 'Asia/Kolkata')::date BETWEEN :a AND :b)
            OR (d.stage NOT IN ('won','lost') AND d.expected_close BETWEEN :a AND :b))
        ORDER BY d.amount_usd DESC NULLS LAST""", acl=acl, bl=business_line_id, owner=owner_id, a=start, b=end)
    undated = _q(s, DEAL_SQL + " AND d.stage NOT IN ('won','lost') AND d.expected_close IS NULL ORDER BY d.amount_usd DESC NULLS LAST",
                 acl=acl, bl=business_line_id, owner=owner_id)
    unconverted = [d["name"] for d in rows + undated if d["value_amount"] is not None and d["amount_usd"] is None]

    groups: Dict[Any, Dict[str, Any]] = {}
    total = _empty_totals()
    deals_out = []
    for d in rows:
        b = _bucket(d)
        amt = _f(d["amount_usd"])
        p = d["probability"] if d["probability"] is not None else probs.get(d["stage"], 0)
        w = amt * p / 100 if b in ("commit", "best_case", "pipeline") else 0.0
        key, label = {"owner": (d["owner_user_id"], d["owner_name"]),
                      "business_line": (d["business_line_id"], d["business_line_name"]),
                      "stage": (d["stage"], STAGE_LABEL.get(d["stage"], d["stage"])),
                      "category": (b, CATEGORY_LABEL.get(b, b))}[group_by]
        g = groups.setdefault(key, {"key": key, "label": label, "t": _empty_totals()})
        for t in (g["t"], total):
            t[b] = t.get(b, 0.0) + amt
            t["weighted"] += w
            t["deals"] += 1
        deals_out.append({"id": d["id"], "name": d["name"], "account_name": d["account_name"], "owner_name": d["owner_name"],
                          "owner_user_id": d["owner_user_id"], "business_line_id": d["business_line_id"],
                          "business_line_name": d["business_line_name"], "stage": d["stage"],
                          "stage_label": STAGE_LABEL.get(d["stage"]), "bucket": b, "category_label": CATEGORY_LABEL.get(b, b),
                          "amount_usd": _f(d["amount_usd"]) if d["amount_usd"] is not None else None,
                          "value_amount": _f(d["value_amount"]) if d["value_amount"] is not None else None,
                          "currency": d["currency"], "probability": p, "probability_is_override": d["probability"] is not None,
                          "weighted_usd": round(w, 2), "close_date": d["closed_on"] if d["stage"] == "won" else d["expected_close"],
                          "next_step": d["next_step"]})
    for t in [total] + [g["t"] for g in groups.values()]:
        t.pop("lost", None)

    # Targets
    owners = visible_owners(s, user)
    owner_ids = [o["id"] for o in owners]
    tg = _q(s, "SELECT user_id, business_line_id, amount_usd FROM sales_targets WHERE period = :p", p=period)
    user_t = {r["user_id"]: _f(r["amount_usd"]) for r in tg if r["user_id"] is not None}
    bl_t = {r["business_line_id"]: _f(r["amount_usd"]) for r in tg if r["business_line_id"] is not None}
    if business_line_id:
        total_target = bl_t.get(business_line_id)
    elif owner_id:
        total_target = user_t.get(owner_id)
    elif group_by == "business_line" or (bl_t and not user_t):
        total_target = sum(bl_t.values()) if bl_t else None
    else:
        vis = [user_t[i] for i in owner_ids if i in user_t]
        total_target = sum(vis) if vis else (sum(bl_t.values()) if bl_t else None)

    def row_target(key):
        if group_by == "owner" and not business_line_id:
            return user_t.get(key)
        if group_by == "business_line" and not owner_id:
            return bl_t.get(key)
        return None

    grid = [{"key": g["key"], "label": g["label"], **_finish(g["t"], row_target(g["key"]))} for g in groups.values()]
    if group_by == "owner" and not business_line_id and not owner_id:          # owners with a target but no deals yet
        for o in owners:
            if o["id"] in user_t and o["id"] not in groups:
                grid.append({"key": o["id"], "label": o["name"], **_finish(_empty_totals(), user_t[o["id"]])})
    order = {"stage": OPEN_STAGES + ["won"], "category": ["closed", "commit", "best_case", "pipeline", "omitted"]}.get(group_by)
    grid.sort(key=(lambda r: order.index(r["key"]) if r["key"] in order else 99) if order
              else (lambda r: -(r["best_case_forecast"] + r["pipeline"])))
    return {"period": period, "start": start, "end": end, "group_by": group_by, "totals": _finish(total, total_target),
            "rows": grid, "deals": deals_out,
            "undated": [{"id": d["id"], "name": d["name"], "account_name": d["account_name"], "owner_name": d["owner_name"],
                         "stage_label": STAGE_LABEL.get(d["stage"]),
                         "amount_usd": _f(d["amount_usd"]) if d["amount_usd"] is not None else None} for d in undated],
            "warnings": ([f"{len(unconverted)} deal(s) use a currency without an exchange rate and count as 0: "
                          + ", ".join(unconverted[:5])] if unconverted else []),
            "stage_probability": probs}


# ── Snapshots & changes ───────────────────────────────────────────────────────


def take_snapshot(s, snap_date: Optional[date] = None) -> int:
    snap_date = snap_date or datetime.now(IST).date()
    n = s.execute(text("""
        INSERT INTO forecast_snapshots (snapshot_date, deal_id, name, account_id, owner_user_id, business_line_id, stage,
                                        forecast_category, amount_usd, probability, expected_close)
        SELECT :d, id, name, account_id, owner_user_id, business_line_id, stage, forecast_category, amount_usd,
               probability, expected_close FROM deals
        ON CONFLICT (snapshot_date, deal_id) DO NOTHING"""), {"d": snap_date}).rowcount
    s.commit()
    return n


def changes(s, user, period: str, since: Optional[str] = None, business_line_id: Optional[int] = None,
            owner_id: Optional[int] = None) -> Dict[str, Any]:
    sm = int(_setting(s, "fiscal_year_start_month", 1))
    start, end = period_bounds(period, sm)
    today = datetime.now(IST).date()
    base = s.execute(text("SELECT max(snapshot_date) FROM forecast_snapshots WHERE snapshot_date <= :d"),
                     {"d": date.fromisoformat(since) if since else today - timedelta(days=1)}).scalar()
    if base is None:
        base = s.execute(text("SELECT min(snapshot_date) FROM forecast_snapshots")).scalar()
    if base is None:
        return {"since": None, "changes": [], "note": "No snapshot yet — the first one is taken automatically."}
    acl = auth.account_scope(s, user)
    cur = {d["id"]: d for d in _q(s, DEAL_SQL, acl=acl, bl=business_line_id, owner=owner_id)}
    old = {d["deal_id"]: d for d in _q(s, """
        SELECT * FROM forecast_snapshots WHERE snapshot_date = :d AND account_id = ANY(:acl)
          AND (CAST(:bl AS int) IS NULL OR business_line_id = :bl) AND (CAST(:owner AS int) IS NULL OR owner_user_id = :owner)""",
        d=base, acl=acl, bl=business_line_id, owner=owner_id)}

    def in_period(close) -> bool:
        return close is not None and start <= close <= end

    out = []
    for did in set(cur) | set(old):
        n, o = cur.get(did), old.get(did)
        n_close = (n["closed_on"] if n and n["stage"] == "won" else n["expected_close"]) if n else None
        o_close = o["expected_close"] if o else None
        if not (in_period(n_close) or in_period(o_close)):
            continue
        name = (n or o)["name"]
        amt = _f(n["amount_usd"]) if n else _f(o["amount_usd"])
        base_row = {"deal_id": did, "name": name, "account_name": n["account_name"] if n else None,
                    "owner_name": n["owner_name"] if n else None, "amount_usd": amt, "exists": n is not None}
        if n and not o:
            out.append({**base_row, "type": "new", "text": "New in the forecast"})
            continue
        if o and not n:
            out.append({**base_row, "type": "removed", "text": "Deleted or moved out of your scope"})
            continue
        if n["stage"] != o["stage"]:
            kind = "won" if n["stage"] == "won" else "lost" if n["stage"] == "lost" else (
                "advanced" if o["stage"] in OPEN_STAGES and n["stage"] in OPEN_STAGES
                and OPEN_STAGES.index(n["stage"]) > OPEN_STAGES.index(o["stage"]) else "moved_back")
            out.append({**base_row, "type": kind, "text": f"{STAGE_LABEL.get(o['stage'])} → {STAGE_LABEL.get(n['stage'])}"})
        if n["forecast_category"] != o["forecast_category"] and n["stage"] not in ("won", "lost"):
            out.append({**base_row, "type": "category", "text": f"{CATEGORY_LABEL.get(o['forecast_category'])} → "
                                                                  f"{CATEGORY_LABEL.get(n['forecast_category'])}"})
        if round(_f(n["amount_usd"]), 2) != round(_f(o["amount_usd"]), 2):
            delta = _f(n["amount_usd"]) - _f(o["amount_usd"])
            out.append({**base_row, "type": "amount", "delta_usd": round(delta, 2),
                        "text": f"Amount {'+' if delta > 0 else '−'}{abs(delta):,.0f} USD"})
        if n["stage"] not in ("won", "lost") and o_close != n["expected_close"]:
            if o_close and n["expected_close"] and n["expected_close"] > o_close:
                kind = "slipped_out" if in_period(o_close) and not in_period(n["expected_close"]) else "slipped"
            elif n["expected_close"] is None:
                kind = "slipped"
            else:
                kind = "pulled_in"
            out.append({**base_row, "type": kind,
                        "text": f"Close date {o_close or 'none'} → {n['expected_close'] or 'none'}"})
    rank = {"won": 0, "lost": 1, "slipped_out": 2, "new": 3, "category": 4, "advanced": 5, "moved_back": 6,
            "slipped": 7, "pulled_in": 8, "amount": 9, "removed": 10}
    out.sort(key=lambda c: (rank.get(c["type"], 99), -c["amount_usd"]))
    return {"since": base, "changes": out}


def _snapshot_loop() -> None:
    """Weekly snapshot every Monday 06:00 IST (+ a baseline the first time)."""
    while True:
        try:
            conn = engine.connect()
            try:
                if conn.execute(text("SELECT pg_try_advisory_lock(84212)")).scalar():
                    now = datetime.now(IST)
                    monday = (now - timedelta(days=now.weekday())).date()
                    s = get_session()
                    try:
                        have_any = s.execute(text("SELECT 1 FROM forecast_snapshots LIMIT 1")).scalar()
                        have_week = s.execute(text("SELECT 1 FROM forecast_snapshots WHERE snapshot_date >= :m LIMIT 1"),
                                              {"m": monday}).scalar()
                        due = now >= datetime.combine(monday, datetime.min.time(), IST).replace(hour=6)
                        if not have_any or (due and not have_week):
                            n = take_snapshot(s, now.date() if not have_any else monday)
                            print(f"[crm] forecast snapshot: {n} deals")
                    finally:
                        s.close()
                    conn.execute(text("SELECT pg_advisory_unlock(84212)"))
            finally:
                conn.close()
        except Exception as e:
            print(f"[crm] forecast snapshot failed: {e}")
        time.sleep(3600)


def start_background() -> None:
    threading.Thread(target=_snapshot_loop, name="crm-forecast-snapshot", daemon=True).start()


# ── Endpoints ─────────────────────────────────────────────────────────────────


@router.get("/meta")
def meta(user: User = Depends(permissions.require_roles(*READERS)), s=Depends(_session)):
    sm = int(_setting(s, "fiscal_year_start_month", 1))
    periods = period_list(sm)
    return {"periods": periods, "current": next(p["key"] for p in periods if p["current"]),
            "fiscal_year_start_month": sm, "owners": visible_owners(s, user),
            "business_lines": _q(s, "SELECT id, name FROM business_lines WHERE active ORDER BY sort, name"),
            "categories": [{"key": k, "label": CATEGORY_LABEL[k]} for k in CATEGORIES],
            "stage_probability": stage_probability(s),
            "fx": [{**r, "usd_rate": _f(r["usd_rate"])} for r in _q(s, "SELECT currency, usd_rate, source, updated_at FROM fx_rates ORDER BY currency")],
            "can_edit_fx": permissions.is_admin(user),
            "can_edit_targets": permissions.is_admin(user) or user.role == "sales_manager",
            "last_snapshot": s.execute(text("SELECT max(snapshot_date) FROM forecast_snapshots")).scalar()}


@router.get("")
def get_forecast(period: str, business_line_id: Optional[int] = None, owner_id: Optional[int] = None,
                 group_by: str = "owner", user: User = Depends(permissions.require_roles(*READERS)), s=Depends(_session)):
    return rollup(s, user, period, business_line_id, owner_id, group_by)


@router.get("/changes")
def get_changes(period: str, since: Optional[str] = None, business_line_id: Optional[int] = None,
                owner_id: Optional[int] = None, user: User = Depends(permissions.require_roles(*READERS)), s=Depends(_session)):
    return changes(s, user, period, since, business_line_id, owner_id)


@router.get("/targets")
def get_targets(period: str, user: User = Depends(permissions.require_roles(*READERS)), s=Depends(_session)):
    period_bounds(period, 1)          # validates the format
    owners = visible_owners(s, user)
    rows = _q(s, """SELECT t.user_id, t.business_line_id, t.amount_usd, coalesce(u.full_name, u.email) AS user_name, bl.name AS bl_name
                    FROM sales_targets t LEFT JOIN users u ON u.id = t.user_id LEFT JOIN business_lines bl ON bl.id = t.business_line_id
                    WHERE t.period = :p""", p=period)
    ut = {r["user_id"]: _f(r["amount_usd"]) for r in rows if r["user_id"]}
    bt = {r["business_line_id"]: _f(r["amount_usd"]) for r in rows if r["business_line_id"]}
    return {"period": period,
            "users": [{"user_id": o["id"], "name": o["name"], "amount_usd": ut.get(o["id"]),
                       "editable": can_edit_target(s, user, o["id"], None)} for o in owners],
            "business_lines": [{"business_line_id": b["id"], "name": b["name"], "amount_usd": bt.get(b["id"]),
                                "editable": can_edit_target(s, user, None, b["id"])}
                               for b in _q(s, "SELECT id, name FROM business_lines WHERE active ORDER BY sort, name")]}


class TargetIn(BaseModel):
    period: str
    user_id: Optional[int] = None
    business_line_id: Optional[int] = None
    amount_usd: Optional[float] = Field(None, ge=0)          # None = remove the target


@router.put("/targets")
def put_target(body: TargetIn, user: User = Depends(permissions.require_write), s=Depends(_session)):
    period_bounds(body.period, 1)
    if (body.user_id is None) == (body.business_line_id is None):
        raise HTTPException(400, "Set a target for either one rep (user_id) or one business line (business_line_id).")
    if not can_edit_target(s, user, body.user_id, body.business_line_id):
        raise HTTPException(403, "You can set targets only for your team" if user.role == "sales_manager" else "Not authorized")
    col, val = ("user_id", body.user_id) if body.user_id is not None else ("business_line_id", body.business_line_id)
    s.execute(text(f"DELETE FROM sales_targets WHERE period = :p AND {col} = :v"), {"p": body.period, "v": val})
    if body.amount_usd is not None:
        s.execute(text(f"""INSERT INTO sales_targets (period, {col}, amount_usd, updated_by) VALUES (:p, :v, :a, :u)"""),
                  {"p": body.period, "v": val, "a": body.amount_usd, "u": user.id})
    s.commit()
    auth.log_audit(s, user.id, "sales_target_set", details=body.model_dump())
    return get_targets(body.period, user, s)


class FxIn(BaseModel):
    currency: str = Field(..., pattern=r"^[A-Z]{3}$")
    usd_rate: float = Field(..., gt=0)


@router.put("/fx")
def put_fx(body: FxIn, user: User = Depends(auth.require_role("super_admin")), s=Depends(_session)):
    if body.currency == "USD":
        raise HTTPException(400, "USD is always 1.")
    s.execute(text("""INSERT INTO fx_rates (currency, usd_rate, source, updated_by, updated_at) VALUES (:c, :r, 'manual', :u, now())
                      ON CONFLICT (currency) DO UPDATE SET usd_rate = EXCLUDED.usd_rate, source = 'manual',
                      updated_by = EXCLUDED.updated_by, updated_at = now()"""), {"c": body.currency, "r": body.usd_rate, "u": user.id})
    n = s.execute(text("UPDATE deals SET currency = currency WHERE upper(currency) = :c"), {"c": body.currency}).rowcount
    s.commit()
    auth.log_audit(s, user.id, "fx_rate_set", details={"currency": body.currency, "usd_rate": body.usd_rate, "deals_recomputed": n})
    return {"currency": body.currency, "usd_rate": body.usd_rate, "deals_recomputed": n}


@router.post("/snapshot")
def snapshot_now(user: User = Depends(auth.require_role("super_admin")), s=Depends(_session)):
    return {"snapshot_date": datetime.now(IST).date(), "deals": take_snapshot(s)}


@router.get("/export")
def export(period: str, business_line_id: Optional[int] = None, owner_id: Optional[int] = None, group_by: str = "owner",
           user: User = Depends(permissions.require_roles(*READERS)), s=Depends(_session)):
    from openpyxl import Workbook
    from apps.sales_copilot import exports
    f = rollup(s, user, period, business_line_id, owner_id, group_by)
    ch = changes(s, user, period, None, business_line_id, owner_id)
    t = f["totals"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    exports._sheet(ws, ["Measure", "USD"], [
        ["Period", f"{period} ({f['start']} – {f['end']})"], ["Target", t["target"]], ["Closed won", t["closed"]],
        ["Commit forecast (closed + commit)", t["commit_forecast"]], ["Best-case forecast", t["best_case_forecast"]],
        ["Open pipeline", t["open_pipeline"]], ["Weighted pipeline", t["weighted"]], ["Gap to target", t["gap"]],
        ["Coverage (open ÷ gap)", t["coverage"]], ["Deals", t["deals"]], ["Exported", exports._now()]], [36, 22])
    cols = ["label", "closed", "commit", "best_case", "pipeline", "weighted", "commit_forecast", "best_case_forecast", "target", "gap", "coverage"]
    exports._sheet(wb.create_sheet("By " + group_by.replace("_", " ")),
                   ["Group", "Closed won", "Commit", "Best case", "Pipeline", "Weighted", "Commit forecast", "Best-case forecast",
                    "Target", "Gap", "Coverage"], [[r[c] for c in cols] for r in f["rows"]], [26] + [14] * 10)
    exports._sheet(wb.create_sheet("Deals"), ["Deal", "Account", "Owner", "Business line", "Stage", "Category", "Amount USD",
                                              "Probability %", "Weighted USD", "Close date", "Next step"],
                   [[d["name"], d["account_name"], d["owner_name"], d["business_line_name"], d["stage_label"], d["category_label"],
                     d["amount_usd"], d["probability"], d["weighted_usd"], str(d["close_date"] or ""), d["next_step"]] for d in f["deals"]],
                   [34, 22, 20, 18, 11, 12, 13, 11, 13, 11, 40])
    exports._sheet(wb.create_sheet("Changes"), ["Change", "Deal", "Account", "Owner", "Detail", "Amount USD"],
                   [[c["type"].replace("_", " "), c["name"], c["account_name"], c["owner_name"], c["text"], c["amount_usd"]]
                    for c in ch["changes"]], [12, 34, 22, 20, 40, 13])
    return Response(exports._xlsx(wb), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{exports.safe_filename("forecast-" + period, "xlsx")}"'})
