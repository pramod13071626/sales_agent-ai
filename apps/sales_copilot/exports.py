"""PDF / Excel / Markdown exports for copilot answers, chats and notes (README §19.3).

PDF: reportlab (already used by pdf_export.py). Excel: openpyxl — real .xlsx with a
styled, frozen, filterable header so it opens cleanly in Excel / Google Sheets.
"""

import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, List, Optional
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = colors.HexColor("#6F6AF8")
MUTED = colors.HexColor("#656D9A")
HEADER_FILL = PatternFill("solid", fgColor="6F6AF8")

_LATIN = {"→": "->", "←": "<-", "“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-", "…": "...",
          "•": "-", "📌": "", "✓": "v", "✕": "x", "≥": ">=", "≤": "<=", "×": "x", " ": " "}


def _latin(s: Any) -> str:
    """reportlab's built-in fonts are Latin-1 only; map common symbols, drop the rest."""
    s = str(s or "")
    for k, v in _LATIN.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "ignore").decode("latin-1")


def _md_to_para(line: str) -> str:
    h = escape(_latin(line))
    h = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", h)
    h = re.sub(r"(^|[^*])\*(?!\s)(.+?)\*(?!\*)", r"\1<i>\2</i>", h)
    h = re.sub(r"(^|\W)_(?!\s)(.+?)_(?=\W|$)", r"\1<i>\2</i>", h)
    h = re.sub(r"\[(\d{1,2})\]", r"<super><font color='#6F6AF8'>[\1]</font></super>", h)
    h = re.sub(r"\[note\]", r"<super><font color='#B8860B'>[your note]</font></super>", h, flags=re.I)
    return h


def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=16, textColor=BRAND, alignment=0, spaceAfter=2),
        "meta": ParagraphStyle("m", parent=ss["Normal"], fontSize=8, textColor=MUTED, spaceAfter=8),
        "q": ParagraphStyle("q", parent=ss["Normal"], fontSize=10, textColor=colors.white, backColor=BRAND,
                            borderPadding=6, leading=13, spaceBefore=6, spaceAfter=8),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=13.5, spaceAfter=4),
        "bullet": ParagraphStyle("bl", parent=ss["Normal"], fontSize=9.5, leading=13.5, leftIndent=12, bulletIndent=2),
        "h": ParagraphStyle("h", parent=ss["Heading3"], fontSize=11, textColor=BRAND, spaceBefore=8, spaceAfter=4),
        "small": ParagraphStyle("s", parent=ss["Normal"], fontSize=7.5, leading=10, textColor=MUTED),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=7.5, leading=9.5),
    }


def _answer_flowables(content: str, st) -> List[Any]:
    out = []
    for raw in (content or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$", line)
        if m:
            out.append(Paragraph(_md_to_para(m.group(1)), st["bullet"], bulletText="•"))
        elif re.match(r"^#{1,4}\s+", line):
            out.append(Paragraph(_md_to_para(re.sub(r"^#{1,4}\s+", "", line)), st["h"]))
        else:
            out.append(Paragraph(_md_to_para(line), st["body"]))
    return out


def _table_flowable(table: Dict[str, Any], st, max_rows: int = 200):
    cols = table.get("columns") or []
    rows = table.get("rows") or []
    if not cols or not rows:
        return None
    data = [[Paragraph(f"<b>{escape(c.replace('_', ' ').title())}</b>", st["cell"]) for c in cols]]
    for r in rows[:max_rows]:
        data.append([Paragraph(escape(_latin(r.get(c, "") if r.get(c) is not None else "")), st["cell"]) for c in cols])
    t = Table(data, repeatRows=1, colWidths=[(180 * mm) / len(cols)] * len(cols))
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEDFE")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E1EC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FB")]),
    ]))
    return t


def _sources_flowables(citations: List[Dict[str, Any]], st) -> List[Any]:
    if not citations:
        return []
    out = [Paragraph("Sources", st["h"])]
    for c in citations:
        date = (c.get("published_at") or "")[:10]
        url = c.get("url") or ""
        link = f' <link href="{escape(url)}" color="#6F6AF8">{escape(_latin(url[:90]))}</link>' if url.startswith("http") else ""
        out.append(Paragraph(f"[{c['n']}] <b>{escape(_latin(c.get('title') or 'Untitled'))}</b> "
                             f"<font color='#656D9A'>({escape(c.get('doc_type') or '')}{', ' + date if date else ''})</font>{link}",
                             st["small"]))
    return out


def _contacts_flowables(contacts: List[Dict[str, Any]], st) -> List[Any]:
    out = []
    for c in contacts or []:
        bits = [f"<b>{escape(_latin(c.get('name')))}</b>", escape(_latin(c.get("title") or "")),
                escape(_latin(c.get("account") or "")), escape(c.get("email") or ""), escape(c.get("phone") or "")]
        out.append(Paragraph("Contact: " + " | ".join(b for b in bits if b), st["small"]))
    return out


def _doc(buf, title: str):
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(15 * mm, 10 * mm, _latin(f"Sales Copilot - {title}"[:110]))
        canvas.drawRightString(195 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()
    return SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                             bottomMargin=18 * mm, title=_latin(title)), footer


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def answer_pdf(question: str, msg: Dict[str, Any], user_name: str) -> bytes:
    st = _styles()
    buf = BytesIO()
    doc, footer = _doc(buf, question[:80] or "Answer")
    extras = msg.get("extras") or {}
    story = [Paragraph("Sales Copilot", st["title"]),
             Paragraph(escape(_latin(f"Prepared for {user_name} - {_now()} - "
                                     f"{'AI answer' if msg.get('mode') == 'llm' else 'From your data'}")), st["meta"]),
             Paragraph(escape(_latin(question)), st["q"])]
    story += _answer_flowables(msg.get("content") or "", st)
    t = _table_flowable(extras.get("table") or {}, st)
    if t:
        story += [Spacer(1, 6), t]
    story += [Spacer(1, 6)] + _contacts_flowables(extras.get("contacts"), st)
    story += _sources_flowables(msg.get("citations") or [], st)
    story += [Spacer(1, 10), Paragraph("Generated from your organisation's sales intelligence data. "
                                       "Verify key facts against the linked sources before sharing externally.", st["small"])]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def chat_pdf(title: str, messages: List[Dict[str, Any]], user_name: str) -> bytes:
    st = _styles()
    buf = BytesIO()
    doc, footer = _doc(buf, title)
    story = [Paragraph(escape(_latin(title or "Copilot chat")), st["title"]),
             Paragraph(escape(_latin(f"Exported by {user_name} - {_now()} - {len(messages)} messages")), st["meta"])]
    for m in messages:
        if m["role"] == "user":
            story.append(Paragraph(escape(_latin(m["content"])), st["q"]))
        else:
            story += _answer_flowables(m["content"], st)
            t = _table_flowable((m.get("extras") or {}).get("table") or {}, st, max_rows=50)
            if t:
                story += [Spacer(1, 4), t]
            story += _sources_flowables(m.get("citations") or [], st)
            story.append(Spacer(1, 8))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _sheet(ws, headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([("" if v is None else (v if isinstance(v, (int, float)) else str(v))) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, 1):
        if widths and i - 1 < len(widths):
            width = widths[i - 1]
        else:
            longest = max([len(str(r[i - 1] or "")) for r in rows[:200]] or [10])
            width = min(60, max(len(h) + 2, longest + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _xlsx(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_xlsx(table: Dict[str, Any], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    cols = table.get("columns") or []
    _sheet(ws, [c.replace("_", " ").title() for c in cols], [[r.get(c) for c in cols] for r in table.get("rows") or []])
    info = wb.create_sheet("About")
    _sheet(info, ["Field", "Value"], [["Question", title], ["Exported", _now()], ["Rows", len(table.get("rows") or [])]], [14, 90])
    return _xlsx(wb)


def sources_xlsx(citations: List[Dict[str, Any]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sources"
    _sheet(ws, ["#", "Title", "Type", "Date", "Link", "Snippet"],
           [[c["n"], c.get("title"), c.get("doc_type"), (c.get("published_at") or "")[:10], c.get("url"), c.get("snippet")]
            for c in citations], [5, 50, 16, 12, 50, 80])
    return _xlsx(wb)


def chat_xlsx(title: str, messages: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Conversation"
    _sheet(ws, ["Time", "Who", "Message", "Answer type"],
           [[str(m.get("created_at") or "")[:19], "You" if m["role"] == "user" else "Copilot", m["content"], m.get("mode") or ""]
            for m in messages], [20, 10, 100, 16])
    src_rows, table_rows, answer_no = [], [], 0
    for m in messages:
        if m["role"] != "assistant":
            continue
        answer_no += 1
        for c in m.get("citations") or []:
            src_rows.append([answer_no, c["n"], c.get("title"), c.get("doc_type"), (c.get("published_at") or "")[:10], c.get("url")])
        tbl = (m.get("extras") or {}).get("table") or {}
        for r in tbl.get("rows") or []:
            table_rows.append([answer_no] + [r.get(c) for c in (tbl.get("columns") or [])])
    _sheet(wb.create_sheet("Sources"), ["Answer #", "#", "Title", "Type", "Date", "Link"], src_rows, [10, 5, 50, 16, 12, 60])
    if table_rows:
        cols = max(len(r) for r in table_rows) - 1
        _sheet(wb.create_sheet("Tables"), ["Answer #"] + [f"Col {i + 1}" for i in range(cols)], table_rows)
    return _xlsx(wb)


def chat_markdown(title: str, messages: List[Dict[str, Any]]) -> bytes:
    lines = [f"# {title or 'Copilot chat'}", "", f"_Exported {_now()}_", ""]
    for m in messages:
        if m["role"] == "user":
            lines += [f"**You:** {m['content']}", ""]
        else:
            lines += [m["content"], ""]
            for c in m.get("citations") or []:
                lines.append(f"[{c['n']}]: {c.get('url') or ''} \"{c.get('title') or ''}\"")
            lines.append("")
    return "\n".join(lines).encode("utf-8")


def notes_xlsx(notes: List[Dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "My notes"
    _sheet(ws, ["Note", "Type", "Person", "Account", "Pinned", "Used", "Created"],
           [[n["text"], n["kind"], n.get("persona"), n.get("account"), "yes" if n.get("pinned") else "",
             n.get("use_count"), str(n.get("created_at") or "")[:19]] for n in notes], [80, 10, 24, 20, 8, 6, 20])
    return _xlsx(wb)


def safe_filename(s: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "-", (s or "copilot").strip())[:60].strip("-") or "copilot"
    return f"{base}.{ext}"
